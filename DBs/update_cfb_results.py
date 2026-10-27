#!/usr/bin/env python3
"""
update_cfb_results.py

Fills PointsFor / PointsAgainst / OT into CFB_2026_Results.xlsx by pulling
final scores from ESPN's public (unofficial, no key required) scoreboard API,
then validates that every game's two mirror rows agree before saving.

Adapted from update_nfl_results.py. Same sheet layout, same CLI, same
"never write a value you aren't confident about" rule. What's different for
college football is called out under CFB DIFFERENCES below.

USAGE
    python update_cfb_results.py --file CFB_2026_Results.xlsx --date 2026-09-26
    python update_cfb_results.py --file CFB_2026_Results.xlsx --start 2026-09-19 --end 2026-09-27
    python update_cfb_results.py --file CFB_2026_Results.xlsx --start 2026-09-19 --end today
    python update_cfb_results.py --file CFB_2026_Results.xlsx --diagnose-date 2026-09-26

If --start/--end/--date are omitted, it scans the whole sheet and fetches every
past date that still has a blank score OR a played game with a blank OT flag.

CFB DIFFERENCES FROM THE NFL SCRIPT
    1. Scoreboard filtering. ESPN's college scoreboard returns only a small
       featured slate unless you pass a conference-group id and a high limit.
       We query group 80 (FBS) and group 81 (FCS) and de-dupe by event id,
       since the sheet includes FCS opponents.
    2. Team matching. The sheet uses Sports-Reference names ("Southern
       California", "Nevada-Las Vegas"); ESPN uses its own ("USC", "UNLV").
       Nicknames can't be used (dozens of Tigers/Bulldogs/Eagles), so teams are
       matched on ESPN's school name after normalising punctuation/accents,
       plus the ALIASES table below. If exactly ONE team in a game resolves,
       the other is inferred from the sheet (a team plays at most one game per
       day, so its row names its opponent) and the learned alias is printed so
       you can add it to ALIASES. If neither resolves, the game is skipped.
    3. OT backfill. The CFB sheet's OT column starts blank (the source CSV had
       no overtime info), so for games that already have scores the script
       verifies ESPN's score matches the sheet and fills in only the OT flag.
       A score mismatch is reported and nothing is written for that game.
    4. Date tolerance. A game is matched to a sheet row within +/- 1 day of
       the date queried, in case ESPN and the sheet disagree about the date of
       a late-night (West Coast / Hawaii) kickoff. Note that only the dates
       you query are fetched, so if a late game stays blank, also run the
       neighbouring date.
    5. Postponed/cancelled games are only used if ESPN marks them completed.
    6. Validation severity. A past-dated game that is still blank is a WARNING
       here, not a blocker (with 230+ teams one unmapped FCS name would
       otherwise stop the whole save). Mirror mismatches, OT mismatches,
       duplicate rows and half-filled games still block the save. Use
       --strict to make blank past games block too, like the NFL script.

IMPORTANT - NOT VERIFIED AGAINST LIVE ESPN
    The group ids, response fields, and every entry in ALIASES were written
    from memory; this script was tested only against a mocked scoreboard.
    Before trusting it, run --diagnose-date on a real Saturday and read the
    output: every team should show a sheet name, and any "[?]" line needs an
    ALIASES entry.
"""

import argparse
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime, timedelta

import openpyxl
import requests

ESPN_SCOREBOARD_URL = "https://site.api.espn.com/apis/site/v2/sports/football/college-football/scoreboard"
GROUPS = ["80", "81"]  # 80 = FBS, 81 = FCS
PAGE_LIMIT = 400
DATE_TOLERANCE_DAYS = 1

# ESPN school name -> name used in the sheet (Sports-Reference style).
# Keys are matched after normalisation (case, accents, punctuation ignored).
# Anything that already equals a sheet name needs no entry. Unverified - extend
# using --diagnose-date.
ALIASES = {
    "USC": "Southern California",
    "LSU": "Louisiana State",
    "TCU": "Texas Christian",
    "SMU": "Southern Methodist",
    "BYU": "Brigham Young",
    "UNLV": "Nevada-Las Vegas",
    "UTEP": "Texas-El Paso",
    "UTSA": "Texas-San Antonio",
    "UAB": "Alabama-Birmingham",
    "UCF": "Central Florida",
    "FIU": "Florida International",
    "Ole Miss": "Mississippi",
    "Pitt": "Pittsburgh",
    "NC State": "North Carolina State",
    "App State": "Appalachian State",
    "Southern Miss": "Southern Mississippi",
    "UConn": "Connecticut",
    "UMass": "Massachusetts",
    "UL Monroe": "Louisiana-Monroe",
    "Miami": "Miami (FL)",          # ESPN's plain "Miami" is the Hurricanes
    "UT Martin": "Tennessee-Martin",
    "UC Davis": "California-Davis",
    "SE Louisiana": "Southeastern Louisiana",
    "Hawai'i": "Hawaii",
    "Sam Houston State": "Sam Houston",
    "McNeese": "McNeese State",
    "Middle Tennessee": "Middle Tennessee State",
    "The Citadel": "Citadel",
    "Grambling": "Grambling State",
    "UAlbany": "Albany",
    "Nicholls": "Nicholls State",
    "Central Connecticut": "Central Connecticut State",
}


def compact(s: str) -> str:
    """Case/accent/punctuation/space-insensitive key: "Hawai'i" -> "hawaii"."""
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", s.lower())


class TeamResolver:
    def __init__(self, sheet_teams):
        self.by_key = {compact(t): t for t in sheet_teams}
        self.aliases = {}
        for espn_name, sheet_name in ALIASES.items():
            if compact(sheet_name) in self.by_key:
                self.aliases[compact(espn_name)] = self.by_key[compact(sheet_name)]
        self.learned = {}

    def resolve(self, team: dict) -> str | None:
        for raw in (team.get("location"), team.get("shortDisplayName")):
            if not raw:
                continue
            key = compact(raw)
            if key in self.aliases:
                return self.aliases[key]
            if key in self.by_key:
                return self.by_key[key]
        return None


def espn_label(team: dict) -> str:
    return team.get("displayName") or team.get("location") or "?"


def fetch_games_for_date(d: date) -> list[dict]:
    """Return one dict per FINAL game on this date (FBS + FCS, de-duped):
    {id, teams: [team_a, team_b], scores: [a, b], ot: 0/1}"""
    games, seen = [], set()
    for group in GROUPS:
        resp = requests.get(
            ESPN_SCOREBOARD_URL,
            params={"dates": d.strftime("%Y%m%d"), "groups": group, "limit": PAGE_LIMIT},
            timeout=20,
        )
        resp.raise_for_status()
        for event in resp.json().get("events", []):
            if event["id"] in seen:
                continue
            seen.add(event["id"])
            status = event.get("status", {})
            stype = status.get("type", {})
            if stype.get("state") != "post" or not stype.get("completed"):
                continue  # not finished / postponed / cancelled
            comp = event["competitions"][0]
            competitors = comp["competitors"]
            if len(competitors) != 2:
                continue
            scores = []
            for c in competitors:
                s = c.get("score")
                scores.append(int(float(s)) if s not in (None, "") else None)
            if None in scores:
                continue
            # 4 quarters in regulation; period > 4 means overtime was played.
            period = status.get("period") or comp.get("status", {}).get("period") or 4
            games.append({
                "id": event["id"],
                "teams": [c["team"] for c in competitors],
                "scores": scores,
                "ot": 1 if period > 4 else 0,
            })
    return games


def diagnose_date(d: date, resolver: TeamResolver) -> None:
    print(f"ESPN teams for {d} (FBS + FCS) and how they map to the sheet:")
    bad = 0
    for g in fetch_games_for_date(d):
        for t in g["teams"]:
            code = resolver.resolve(t)
            flag = "   " if code else "[?]"
            if not code:
                bad += 1
            print(f"  {flag} {espn_label(t)!r:42s} location={t.get('location')!r:28s} -> {code}")
    print(f"\n{bad} unmapped team(s). (A '[?]' team is fine if it simply isn't on the sheet; "
          f"otherwise add it to ALIASES.)")


def to_date(cell_value):
    return cell_value.date() if hasattr(cell_value, "date") else cell_value


def col_index(ws, header_name: str) -> int:
    for cell in ws[1]:
        if cell.value == header_name:
            return cell.column
    raise KeyError(f"Column {header_name!r} not found in header row")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--file", required=True, help="Path to CFB_2026_Results.xlsx")
    ap.add_argument("--start", help="YYYY-MM-DD, or 'today'")
    ap.add_argument("--end", help="YYYY-MM-DD, or 'today'")
    ap.add_argument("--date", help="Single date YYYY-MM-DD, shorthand for --start/--end the same day")
    ap.add_argument("--diagnose-date", help="Print ESPN team names + mapping for YYYY-MM-DD and exit, no file changes")
    ap.add_argument("--strict", action="store_true", help="Treat past-dated blank games as blocking errors")
    ap.add_argument("--skip-ot-backfill", action="store_true",
                    help="Don't refetch dates just to fill OT on games that already have scores")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.file)
    ws = wb.active

    c_date = col_index(ws, "Date")
    c_team = col_index(ws, "Team")
    c_opp = col_index(ws, "Opp")
    c_pf = col_index(ws, "PointsFor")
    c_pa = col_index(ws, "PointsAgainst")
    c_ot = col_index(ws, "OT")

    rows = [r for r in ws.iter_rows(min_row=2) if r[c_date - 1].value is not None]
    resolver = TeamResolver({r[c_team - 1].value for r in rows} | {r[c_opp - 1].value for r in rows})

    if args.diagnose_date:
        diagnose_date(datetime.strptime(args.diagnose_date, "%Y-%m-%d").date(), resolver)
        return

    today = date.today()

    # Lookups: exact (date, team, opp) -> row, and team -> [(date, row)] for tolerant matching.
    lookup = {}
    by_team = defaultdict(list)
    for r in rows:
        gdate = to_date(r[c_date - 1].value)
        team, opp = r[c_team - 1].value, r[c_opp - 1].value
        lookup[(gdate, team, opp)] = r
        by_team[team].append((gdate, r))

    def rows_near(team, d):
        return [(gd, r) for gd, r in by_team[team] if abs((gd - d).days) <= DATE_TOLERANCE_DAYS]

    def find_row(team, opp, d):
        hits = [(gd, r) for gd, r in rows_near(team, d) if r[c_opp - 1].value == opp]
        hits.sort(key=lambda x: abs((x[0] - d).days))  # prefer exact date
        return hits[0][1] if hits else None

    # Figure out which dates to fetch.
    if args.date:
        target_dates = {datetime.strptime(args.date, "%Y-%m-%d").date()}
    elif args.start or args.end:
        start = today if args.start in (None, "today") else datetime.strptime(args.start, "%Y-%m-%d").date()
        end = today if args.end in (None, "today") else datetime.strptime(args.end, "%Y-%m-%d").date()
        target_dates = {start + timedelta(days=i) for i in range((end - start).days + 1)}
    else:
        target_dates = set()
        for r in rows:
            gdate = to_date(r[c_date - 1].value)
            pf, ot = r[c_pf - 1].value, r[c_ot - 1].value
            needs_score = pf is None
            needs_ot = (pf is not None and ot is None and not args.skip_ot_backfill)
            if gdate <= today and (needs_score or needs_ot):
                target_dates.add(gdate)

    target_dates = sorted(d for d in target_dates if d <= today)
    if not target_dates:
        print("No past dates to update.")
    else:
        print(f"Fetching {len(target_dates)} date(s): {target_dates[0]} to {target_dates[-1]}")

    filled = ot_filled = already_complete = unmatched = conflicts = 0

    for d in target_dates:
        try:
            games = fetch_games_for_date(d)
        except requests.RequestException as e:
            print(f"  [!] Failed to fetch {d}: {e}")
            continue

        for g in games:
            ta, tb = g["teams"]
            sa, sb = g["scores"]
            ca, cb = resolver.resolve(ta), resolver.resolve(tb)
            label = f"{espn_label(ta)} vs {espn_label(tb)}"

            if ca is None and cb is None:
                print(f"  [!] {d}: could not map either team ({label}) - skipped. "
                      f"Not on sheet, or add to ALIASES.")
                unmatched += 1
                continue
            if ca is None or cb is None:
                known, unknown_team = (ca, tb) if ca else (cb, ta)
                near = rows_near(known, d)
                if len(near) != 1:
                    print(f"  [!] {d}: could not map one team in {label} and the sheet doesn't "
                          f"pin down {known}'s opponent - skipped. Add to ALIASES.")
                    unmatched += 1
                    continue
                inferred = near[0][1][c_opp - 1].value
                print(f"  [~] {d}: inferred ESPN {espn_label(unknown_team)!r} (location "
                      f"{unknown_team.get('location')!r}) = {inferred!r} from {known}'s schedule "
                      f"row - add to ALIASES to make this permanent")
                if ca is None:
                    ca = inferred
                else:
                    cb = inferred

            row_a, row_b = find_row(ca, cb, d), find_row(cb, ca, d)
            if row_a is None or row_b is None:
                print(f"  [!] {d}: no matching sheet rows for {ca} vs {cb} - game not on sheet?")
                unmatched += 1
                continue

            a_pf, b_pf = row_a[c_pf - 1].value, row_b[c_pf - 1].value

            if a_pf is None and b_pf is None:
                for row, pf, pa in ((row_a, sa, sb), (row_b, sb, sa)):
                    row[c_pf - 1].value = pf
                    row[c_pa - 1].value = pa
                    row[c_ot - 1].value = g["ot"]
                filled += 1
            elif a_pf is not None and b_pf is not None:
                same = (a_pf == sa and row_a[c_pa - 1].value == sb and
                        b_pf == sb and row_b[c_pa - 1].value == sa)
                if not same:
                    print(f"  [!] {d}: {ca} vs {cb}: sheet score ({a_pf}-{row_a[c_pa - 1].value}) "
                          f"differs from ESPN ({sa}-{sb}) - left untouched")
                    conflicts += 1
                elif row_a[c_ot - 1].value is None and row_b[c_ot - 1].value is None:
                    row_a[c_ot - 1].value = row_b[c_ot - 1].value = g["ot"]
                    ot_filled += 1
                else:
                    already_complete += 1
            else:
                print(f"  [!] {d}: {ca} vs {cb}: only ONE of the two rows has a score - left untouched")
                conflicts += 1

    print(f"Scores filled for {filled} game(s); OT backfilled on {ot_filled} already-scored game(s); "
          f"{already_complete} already complete; {conflicts} conflict(s); {unmatched} unmatched.")

    # ---- Validation pass over the WHOLE sheet before saving ----
    errors, warnings = [], []
    seen = set()
    ot_unknown = 0
    for r in rows:
        gdate = to_date(r[c_date - 1].value)
        team, opp = r[c_team - 1].value, r[c_opp - 1].value
        pf, pa, ot = r[c_pf - 1].value, r[c_pa - 1].value, r[c_ot - 1].value

        dupe_key = (gdate, team, opp)
        if dupe_key in seen:
            errors.append(f"{gdate} {team} vs {opp}: duplicate row")
        seen.add(dupe_key)

        if pf is None:
            if gdate < today:
                warnings.append(f"{gdate} {team} vs {opp}: game is in the past but still has no score")
            continue  # future game, fine

        mirror = lookup.get((gdate, opp, team))
        if mirror is None:
            errors.append(f"{gdate} {team} vs {opp}: no mirror row found at all")
            continue
        m_pf, m_pa, m_ot = mirror[c_pf - 1].value, mirror[c_pa - 1].value, mirror[c_ot - 1].value
        if m_pf is None:
            errors.append(f"{gdate} {team} vs {opp}: mirror row ({opp} vs {team}) has no score")
            continue
        if pf != m_pa or pa != m_pf:
            errors.append(f"{gdate} {team}({pf}-{pa}) vs mirror {opp}({m_pf}-{m_pa}): scores don't mirror")
        if ot != m_ot:
            errors.append(f"{gdate} {team} vs {opp}: OT flag mismatch ({ot} vs {m_ot})")
        if ot is None:
            ot_unknown += 1

    if ot_unknown:
        print(f"Note: {ot_unknown // 2} scored game(s) still have a blank OT flag.")
    if warnings:
        print(f"\n{len(warnings) // 2 or len(warnings)} past game(s) still blank "
              f"({len(warnings)} rows){' - BLOCKING (--strict)' if args.strict else ''}:")
        for w in warnings[:40]:
            print(f"  [~] {w}")
        if len(warnings) > 40:
            print(f"  ... and {len(warnings) - 40} more rows")
        if args.strict:
            errors.extend(warnings)

    if errors:
        print(f"\n{len(errors)} validation issue(s) found - NOT saving the file:")
        for e in errors:
            print(f"  [!] {e}")
        sys.exit(1)

    wb.save(args.file)
    print(f"\nSaved {args.file} - sheet passed validation.")


if __name__ == "__main__":
    main()
