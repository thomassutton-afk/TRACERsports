#!/usr/bin/env python3
"""
Convert a Basketball-Reference NBA schedule CSV into the TRACER "NBA_<year>_Results.xlsx" layout.

    python build_nba_results.py NBA_2027_Schedule.csv
    python build_nba_results.py NBA_2027_Schedule.csv -o NBA_2027_Results.xlsx
    python build_nba_results.py NBA_2027_Schedule.csv --season 2027

Output layout (matches the 2026 workbook):
  * One sheet, "Sheet1", one Excel Table ("Table1", TableStyleLight8).
  * Columns: Date, Season, Type, Round, Team, Opp, HomeAway, PF, PA, OT
  * Two rows per game, in this order: visitor row (A), then home row (H).
    Neutral-site games get N on both rows.
  * Season = the year the season ENDS (Oct 2026 - Jun 2027 -> 2027).
  * Type/Round = R / RS for every row (the schedule CSV is regular season only).
  * PF/PA/OT are left blank for games with no score yet. Re-run the script on an
    updated CSV (one that has PTS filled in) and they populate. OT is 1 if the CSV
    marks the game as going to overtime (OT, 2OT, ...), else 0.

Things the script cleans up in the source CSV:
  * Basketball-Reference's citation line / blank lines / repeated header rows.
  * Games that appear twice (exact duplicate blocks).
"""
import argparse
import csv
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles.colors import Color
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.utils.indexed_list import IndexedList
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------------------------------- #
# Config
# --------------------------------------------------------------------------- #
AUTHOR = "Tj Sutton"

# Abbreviations as used in the existing results workbook (note BRK and CHH).
TEAM_ABBR = {
    "Atlanta Hawks": "ATL",
    "Boston Celtics": "BOS",
    "Brooklyn Nets": "BRK",
    "Charlotte Hornets": "CHH",
    "Chicago Bulls": "CHI",
    "Cleveland Cavaliers": "CLE",
    "Dallas Mavericks": "DAL",
    "Denver Nuggets": "DEN",
    "Detroit Pistons": "DET",
    "Golden State Warriors": "GSW",
    "Houston Rockets": "HOU",
    "Indiana Pacers": "IND",
    "Los Angeles Clippers": "LAC",
    "Los Angeles Lakers": "LAL",
    "Memphis Grizzlies": "MEM",
    "Miami Heat": "MIA",
    "Milwaukee Bucks": "MIL",
    "Minnesota Timberwolves": "MIN",
    "New Orleans Pelicans": "NOP",
    "New York Knicks": "NYK",
    "Oklahoma City Thunder": "OKC",
    "Orlando Magic": "ORL",
    "Philadelphia 76ers": "PHI",
    "Phoenix Suns": "PHX",
    "Portland Trail Blazers": "POR",
    "Sacramento Kings": "SAC",
    "San Antonio Spurs": "SAS",
    "Toronto Raptors": "TOR",
    "Utah Jazz": "UTA",
    "Washington Wizards": "WAS",
}

# A game is marked N (neutral) when it is played somewhere other than the home
# team's regular arena (international games, etc.). A blank Arena in the CSV is
# treated the same way (venue not confirmed as the home arena).
# Arenas listed here are "alternate home courts" that still count as H.
ALTERNATE_HOME_ARENAS = {"Moody Center"}  # Spurs home games in Austin

# Manual overrides: {("YYYY-MM-DD", "VIS", "HOME"): "N" or "H"}
HOMEAWAY_OVERRIDES = {}

# --------------------------------------------------------------------------- #
# Workbook formatting constants (copied from the 2026 workbook)
# --------------------------------------------------------------------------- #
SHEET_NAME = "Sheet1"
HEADERS = ["Date", "Season", "Type", "Round", "Team", "Opp", "HomeAway", "PF", "PA", "OT"]
DATE_FORMAT = "mm-dd-yy"  # builtin numFmtId 14 (displays as m/d/yyyy in Excel)
COL_WIDTHS = {
    "A": 9.9140625,
    "C": 7.08203125,
    "D": 8.5,
    "E": 7.58203125,
    "F": 6.58203125,
    "G": 12.75,
    "H": 5.33203125,
    "I": 5.5,
    "J": 5.6640625,
}  # column B keeps the default width, as in the original
CSV_DATE_FORMAT = "%a %b %d %Y"  # e.g. "Tue Oct 20 2026"


# --------------------------------------------------------------------------- #
# Reading the CSV
# --------------------------------------------------------------------------- #
def to_int(text):
    text = (text or "").strip()
    return int(text) if text.isdigit() else None


def read_games(path):
    with open(path, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))

    header = next((r for r in rows if r and r[0].strip() == "Date"), None)
    if header is None:
        sys.exit("Could not find the 'Date' header row in the CSV.")
    try:
        i_vis = header.index("Visitor/Neutral")
        i_home = header.index("Home/Neutral")
        i_arena = header.index("Arena")
    except ValueError as e:
        sys.exit(f"Unexpected CSV header: {e}")
    i_vpts, i_hpts = i_vis + 1, i_home + 1
    # Basketball-Reference exports two unnamed columns after the home PTS:
    # box-score link, then OT. Use a named OT column if one ever exists.
    i_ot = header.index("OT") if "OT" in header else i_hpts + 2

    games, order, skipped, dupes = {}, [], 0, 0
    for r in rows:
        if not r:
            continue
        try:
            date = datetime.strptime(r[0].strip(), CSV_DATE_FORMAT)
        except ValueError:
            skipped += 1  # citation line, blank lines, repeated header rows
            continue
        g = {
            "date": date,
            "vis": r[i_vis].strip(),
            "home": r[i_home].strip(),
            "vpts": to_int(r[i_vpts]),
            "hpts": to_int(r[i_hpts]),
            "ot": bool(r[i_ot].strip()) if len(r) > i_ot else False,
            "arena": r[i_arena].strip() if len(r) > i_arena else "",
        }
        key = (g["date"], g["vis"], g["home"])
        if key in games:
            dupes += 1
            # keep the copy that has a score if only one does
            if games[key]["vpts"] is None and g["vpts"] is not None:
                games[key] = g
            continue
        games[key] = g
        order.append(key)

    out = [games[k] for k in order]
    out.sort(key=lambda g: g["date"])  # stable: same-day games keep CSV order
    return out, skipped, dupes


# --------------------------------------------------------------------------- #
# Transforming
# --------------------------------------------------------------------------- #
def season_of(d):
    return d.year + 1 if d.month >= 8 else d.year


def classify_venues(games):
    """Return {id(game): 'H' or 'N'} and a list of alternate-venue games to report."""
    arenas = defaultdict(Counter)
    for g in games:
        if g["arena"]:
            arenas[g["home"]][g["arena"]] += 1
    primary = {t: c.most_common(1)[0][0] for t, c in arenas.items()}

    result, report = {}, []
    for g in games:
        flag = "H"
        if g["arena"] != primary.get(g["home"]) and g["arena"] not in ALTERNATE_HOME_ARENAS:
            flag = "N"
        key = (g["date"].strftime("%Y-%m-%d"), TEAM_ABBR.get(g["vis"]), TEAM_ABBR.get(g["home"]))
        if key in HOMEAWAY_OVERRIDES:
            flag = HOMEAWAY_OVERRIDES[key]
        result[id(g)] = flag
        if g["arena"] != primary.get(g["home"]):
            report.append((g, flag))
    return result, report


def build_rows(games, season_override):
    unknown = {n for g in games for n in (g["vis"], g["home"]) if n not in TEAM_ABBR}
    if unknown:
        sys.exit(f"Unknown team name(s) in CSV (add to TEAM_ABBR): {sorted(unknown)}")

    venue, report = classify_venues(games)
    rows = []
    for g in games:
        v, h = TEAM_ABBR[g["vis"]], TEAM_ABBR[g["home"]]
        played = g["vpts"] is not None and g["hpts"] is not None
        ot = (1 if g["ot"] else 0) if played else None
        season = season_override or season_of(g["date"])
        v_ha, h_ha = ("N", "N") if venue[id(g)] == "N" else ("A", "H")
        d = g["date"]
        rows.append([d, season, "R", "RS", v, h, v_ha, g["vpts"], g["hpts"], ot])
        rows.append([d, season, "R", "RS", h, v, h_ha, g["hpts"], g["vpts"], ot])
    return rows, report


# --------------------------------------------------------------------------- #
# Writing the workbook
# --------------------------------------------------------------------------- #
def write_workbook(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    # Workbook default font: Times New Roman 12 (every cell inherits it)
    wb._fonts = IndexedList(
        [Font(name="Times New Roman", sz=12, family=2, color=Color(theme=1))]
    )
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR

    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    for (cell,) in ws.iter_rows(min_row=2, max_col=1):
        cell.number_format = DATE_FORMAT

    for letter, width in COL_WIDTHS.items():
        dim = ws.column_dimensions[letter]
        dim.width = width
        dim.bestFit = True
    ws.sheet_format.defaultRowHeight = 15.5
    ws.page_margins = PageMargins(left=0.7, right=0.7, top=0.75, bottom=0.75, header=0.3, footer=0.3)
    ws.sheet_view.tabSelected = True

    # Excel Table, same name/style as the original; Date column carries the
    # m/d/yyyy data-format (dxf 0) just like the 2026 workbook.
    table = Table(displayName="Table1", ref=f"A1:J{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleLight8",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table._initialise_columns()
    for col, name in zip(table.tableColumns, HEADERS):
        col.name = name
    wb._differential_styles.add(
        DifferentialStyle(numFmt=NumberFormat(numFmtId=19, formatCode="m/d/yyyy"))
    )
    table.tableColumns[0].dataDxfId = 0
    ws.add_table(table)

    wb.save(out_path)


# --------------------------------------------------------------------------- #
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("csv", type=Path, help="Basketball-Reference schedule CSV")
    ap.add_argument("-o", "--output", type=Path, help="output .xlsx (default: *_Schedule.csv -> *_Results.xlsx)")
    ap.add_argument("--season", type=int, help="force the Season value (default: year the season ends)")
    args = ap.parse_args()

    out = args.output or args.csv.with_name(
        args.csv.stem.replace("Schedule", "Results") + ".xlsx"
        if "Schedule" in args.csv.stem
        else args.csv.stem + "_Results.xlsx"
    )

    games, skipped, dupes = read_games(args.csv)
    rows, report = build_rows(games, args.season)
    write_workbook(rows, out)

    played = sum(1 for g in games if g["vpts"] is not None and g["hpts"] is not None)
    per_team = Counter()
    for g in games:
        per_team[g["vis"]] += 1
        per_team[g["home"]] += 1
    seasons = sorted({r[1] for r in rows})

    print(f"Wrote {out}")
    print(f"  {len(games)} games -> {len(rows)} rows | season(s): {seasons} | "
          f"{played} with scores, {len(games) - played} unplayed")
    print(f"  cleaned: {skipped} non-data lines skipped, {dupes} duplicate games dropped")
    print(f"  games per team: min {min(per_team.values())}, max {max(per_team.values())} "
          f"(a full season is 82)")
    if report:
        print("  Non-home-arena games (please eyeball these):")
        for g, flag in report:
            print(f"    {g['date']:%Y-%m-%d} {TEAM_ABBR[g['vis']]}@{TEAM_ABBR[g['home']]}  "
                  f"arena={g['arena'] or '(blank)'}  -> {flag}")


if __name__ == "__main__":
    main()
