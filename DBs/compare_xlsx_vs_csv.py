"""Compare scores in CFB_2026_Results.xlsx against the Sports-Reference CSV.
Usage (from DBs):  python compare_xlsx_vs_csv.py cfb\\Results\\CFB_2026_Results_test.xlsx cfb\\Results\\CFB_2026_Results.csv
"""
import csv, re, sys, datetime
import openpyxl

xlsx_path, csv_path = sys.argv[1], sys.argv[2]
rank = lambda t: re.sub(r"^\(\d+\)\s+", "", t.strip())

# CSV truth: {(date, frozenset({winner, loser})): (winner, winner_pts, loser_pts)}
truth = {}
for r in list(csv.reader(open(csv_path, encoding="utf-8", errors="replace")))[5:]:
    if len(r) < 10 or not r[2].strip() or not (r[6].strip() and r[9].strip()):
        continue
    d = datetime.datetime.strptime(r[2].strip(), "%b %d %Y").date()
    w, l = rank(r[5]), rank(r[8])
    truth[(d, frozenset((w, l)))] = (w, int(r[6]), int(r[9]))

# xlsx: two mirror rows per game
ws = openpyxl.load_workbook(xlsx_path).worksheets[0]
hdr = [c.value for c in ws[1]]
ix = {h: hdr.index(h) for h in ("Date", "Team", "Opp", "PointsFor", "PointsAgainst", "OT")}
got, ot_filled = {}, 0
for row in ws.iter_rows(min_row=2, values_only=True):
    if row[ix["PointsFor"]] is None:
        continue
    d = row[ix["Date"]].date()
    got[(d, row[ix["Team"]], row[ix["Opp"]])] = (row[ix["PointsFor"]], row[ix["PointsAgainst"]])
    ot_filled += row[ix["OT"]] not in (None, "")

missing, bad = [], []
for (d, pair), (w, wp, lp) in truth.items():
    l = next(t for t in pair if t != w)
    g = got.get((d, w, l))
    if g is None:
        missing.append((d, w, l))
    elif g != (wp, lp):
        bad.append((d, w, l, g, (wp, lp)))
truth_keys = {(d, w, l) for (d, pair), (w, _, _) in truth.items() for l in pair if l != w}
extra = [k for k in got if k not in truth_keys and (k[0], k[2], k[1]) not in truth_keys]
print(f"CSV scored games: {len(truth)} | xlsx scored rows: {len(got)} | OT cells filled: {ot_filled}")
print(f"missing from xlsx: {len(missing)} | score mismatches: {len(bad)} | in xlsx but not CSV: {len(extra)}")
for m in missing[:15]: print("  missing ", m)
for b in bad[:15]: print("  mismatch", b)
for e in extra[:15]: print("  extra   ", e)
