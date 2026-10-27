"""Apply the display names in cfb_team_names.xlsx to cfb_elo.db.

Run from DBs/cfb (same place as add_season.py):

    python apply_display_names.py cfb_team_names.xlsx            # dry run: shows what would change
    python apply_display_names.py cfb_team_names.xlsx --apply    # writes to the DB + the json

What it writes (the same two places add_season.py's docstring says to use):
  * teams.team_name              <- 'Display full name'
  * team_history.name            <- 'Display full name', on each team's currently-open era only
                                    (past eras keep their own names, e.g. Dixie State)
  * team_display_names.json      <- {slug: {team_id, full, short, sports_reference, espn}} for
                                    the site config; the short names live here, not in the DB.

It never touches slugs/codes, games, ratings or schedule. Safe to re-run.
"""
import argparse, json, os, sqlite3, sys
from collections import defaultdict
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import db  # DBs/cfb/db.py

FULL, SHORT = "Display full name", "Display short name"

def read_sheet(path):
    ws = load_workbook(path, data_only=True)["Teams"]
    head = [c.value for c in ws[1]]
    need = ["Slug", "Team ID", "Sports-Reference name", "ESPN full name", FULL, SHORT]
    missing = [h for h in need if h not in head]
    if missing:
        sys.exit(f"Teams sheet is missing column(s): {missing}")
    return [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--db", default="../cfb_elo.db")
    ap.add_argument("--json-out", default="team_display_names.json")
    ap.add_argument("--apply", action="store_true", help="write changes (default is a dry run)")
    a = ap.parse_args()

    rows = read_sheet(a.xlsx)
    conn = db.connect(a.db)
    teams = {tid: name for tid, name in conn.execute("SELECT team_id, team_name FROM teams")}
    slugs = {tid: code for tid, code in conn.execute(
        "SELECT team_id, code FROM team_history WHERE end_season IS NULL")}

    problems, warnings = [], []
    seen_ids = set()
    for r in rows:
        tid, slug = r["Team ID"], r["Slug"]
        full, short = (r[FULL] or "").strip(), (r[SHORT] or "").strip()
        if tid not in teams:
            problems.append(f"{slug}: team_id {tid} is not in the DB"); continue
        if slugs.get(tid) not in (None, slug):
            problems.append(f"{slug}: sheet slug does not match the DB code {slugs[tid]} for {tid}")
        if not full: problems.append(f"{slug}: blank '{FULL}'")
        if not short: problems.append(f"{slug}: blank '{SHORT}'")
        seen_ids.add(tid)
    for tid in set(teams) - seen_ids:
        warnings.append(f"{teams[tid]} ({tid}) is in the DB but not in the sheet - left unchanged")
    for col in (FULL, SHORT):
        by = defaultdict(list)
        for r in rows: by[(r[col] or "").strip().lower()].append(r["Slug"])
        for name, ss in by.items():
            if name and len(ss) > 1:
                warnings.append(f"duplicate {col.lower()} '{name}': {', '.join(ss)} "
                                f"(two team_ids for one school? see franchise.py)")

    if problems:
        print("PROBLEMS (nothing written):")
        for p in problems: print("  -", p)
        sys.exit(1)

    changes = [(r["Team ID"], r["Slug"], teams[r["Team ID"]], r[FULL].strip()) for r in rows
               if teams[r["Team ID"]] != r[FULL].strip()]
    print(f"{len(rows)} teams in sheet | {len(changes)} DB name(s) would change | "
          f"{len(rows) - len(changes)} already match")
    for tid, slug, old, new in changes[:8]:
        print(f"    {tid}  {old!r} -> {new!r}")
    if len(changes) > 8: print(f"    ... and {len(changes) - 8} more")
    for w in warnings: print("  warning:", w)

    payload = {r["Slug"]: dict(team_id=r["Team ID"], full=r[FULL].strip(), short=r[SHORT].strip(),
                               sports_reference=r["Sports-Reference name"], espn=r["ESPN full name"])
               for r in rows}

    if not a.apply:
        print("\nDry run only. Re-run with --apply to write (back up cfb_elo.db first).")
        return

    no_era = []
    for tid, slug, old, new in changes:
        db.upsert_team(conn, tid, new)
        if not db.rename_current_history(conn, tid, new):
            no_era.append(slug)
    conn.commit()
    with open(a.json_out, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=1, sort_keys=True)
    print(f"\nApplied {len(changes)} name(s); wrote {a.json_out} ({len(payload)} teams).")
    if no_era:
        print(f"  {len(no_era)} team(s) have no open team_history row, so only teams.team_name "
              f"changed: {', '.join(no_era[:10])}")

if __name__ == "__main__":
    main()
