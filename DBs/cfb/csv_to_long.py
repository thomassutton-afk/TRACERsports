#!/usr/bin/env python3
"""
Convert a Sports-Reference-style results CSV (one row per game) into the
long format used by NFL_2026_Results.xlsx (two rows per game, one from each
team's perspective).

Usage:
    python csv_to_long.py CFB_2026_Results.csv CFB_2026_Results.xlsx

Input columns (after the junk header lines):
    Rk, Wk, Date, Time, Day, Winner, Pts, <loc>, Loser, Pts, Notes

The unnamed <loc> column sits between the two Pts columns:
    ''  -> game at the Winner's site
    '@' -> game at the Loser's site
    'N' -> neutral site
For unplayed games the points are blank and the first-listed team is the
visitor (so '@' still means "second team is home").

Output columns:
    Date, Season, Type, Round, Team, Opp, HomeAway, PointsFor, PointsAgainst, OT
"""
import re
import sys
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font

COLS = ["Date", "Season", "Type", "Round", "Team", "Opp",
        "HomeAway", "PointsFor", "PointsAgainst", "OT"]

RANK_RE = re.compile(r"^\(\d+\)\s+")  # "(14) Southern California" -> "Southern California"


def read_source(path):
    # The file starts with a citation line and blank lines; find the real header.
    with open(path, encoding="utf-8-sig") as f:
        lines = f.read().splitlines()
    header_idx = next(i for i, ln in enumerate(lines) if ln.startswith("Rk,"))
    df = pd.read_csv(path, skiprows=header_idx, dtype=str, keep_default_na=False)
    df.columns = ["Rk", "Wk", "Date", "Time", "Day",
                  "Winner", "WPts", "Loc", "Loser", "LPts", "Notes"]
    df = df[df["Rk"].str.fullmatch(r"\d+")]  # drop any repeated header / blank rows
    return df.reset_index(drop=True)


def to_int(s):
    return int(s) if s != "" else None


def convert(df):
    rows = []
    for r in df.itertuples(index=False):
        date = datetime.strptime(r.Date, "%b %d %Y")
        season = date.year if date.month >= 6 else date.year - 1
        first = RANK_RE.sub("", r.Winner).strip()   # winner (or visitor if unplayed)
        second = RANK_RE.sub("", r.Loser).strip()   # loser (or host if unplayed)
        p1, p2 = to_int(r.WPts), to_int(r.LPts)

        if r.Loc == "@":        # second-listed team is home
            home, away, hp, ap, ha_home, ha_away = second, first, p2, p1, "H", "A"
        elif r.Loc == "N":      # neutral: keep source order
            home, away, hp, ap, ha_home, ha_away = first, second, p1, p2, "N", "N"
        else:                   # blank: first-listed team is home
            home, away, hp, ap, ha_home, ha_away = first, second, p1, p2, "H", "A"

        base = dict(Date=date, Season=season, Type="R", Round="RS", OT=None)
        rows.append({**base, "Team": home, "Opp": away, "HomeAway": ha_home,
                     "PointsFor": hp, "PointsAgainst": ap})
        rows.append({**base, "Team": away, "Opp": home, "HomeAway": ha_away,
                     "PointsFor": ap, "PointsAgainst": hp})
    return pd.DataFrame(rows, columns=COLS)


def write_xlsx(out, path):
    wb = Workbook()
    ws = wb.active
    ws.title = path.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    font = Font(name="Times New Roman", size=12)
    ws.append(COLS)
    for rec in out.itertuples(index=False):
        ws.append([None if pd.isna(v) else (v.to_pydatetime() if hasattr(v, "to_pydatetime") else v)
                   for v in rec])
    for row in ws.iter_rows():
        for c in row:
            c.font = font
    for c in ws["A"][1:]:
        c.number_format = "mm-dd-yy"
    wb.save(path)


if __name__ == "__main__":
    src, dst = sys.argv[1], sys.argv[2]
    out = convert(read_source(src))
    write_xlsx(out, dst)
    print(f"{len(out)} rows ({len(out)//2} games) -> {dst}")
